"""Ręczne poprawki nazw miejscowości stacji trafo (wspólne dla Rysunków stałych, PZT, opisów).

Plik ``nazwy miejscowości.csv`` (w folderze projektu): ``numer;miejscowość`` - tylko stacje
poprawione ręcznie. Poprawka ma pierwszeństwo przed geokodowaniem i Warunkami Technicznymi.
Zapis może od razu poprawić też ``zestawienie_stacji.xlsx`` (Opisy/PPW/Zakładki czytają nazwy
z niego), bez ponownego generowania rysunków stałych.
"""
from __future__ import annotations

import csv
import os
import re

NAZWA_PLIKU = "nazwy miejscowości.csv"


def nr_stacji(tekst) -> int | None:
    """„03-0541” / „0541” / „541” -> 541."""
    m = re.search(r"(\d+)\s*$", str(tekst or "").strip())
    return int(m.group(1)) if m else None


def wczytaj(sciezka: str) -> dict:
    """{nr stacji (int): miejscowość} albo {} gdy brak pliku."""
    try:
        with open(sciezka, encoding="utf-8-sig", newline="") as f:
            tekst = f.read()
    except OSError:
        return {}
    wynik = {}
    for wiersz in csv.reader(tekst.splitlines(), delimiter=";"):
        if len(wiersz) >= 2 and wiersz[1].strip():
            nr = nr_stacji(wiersz[0])
            if nr is not None:
                wynik[nr] = wiersz[1].strip()
    return wynik


def zapisz(sciezka: str, nazwy: dict) -> None:
    """Zapisuje poprawki ({nr: miejscowość}); pusty słownik usuwa plik."""
    if not nazwy:
        if os.path.exists(sciezka):
            os.remove(sciezka)
        return
    with open(sciezka, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["numer", "miejscowosc"])
        for nr, nazwa in sorted(nazwy.items()):
            w.writerow([nr, nazwa])


def z_zestawienia(xlsx: str) -> list:
    """[(oznaczenie stacji, nr, miejscowość)] z zestawienia stacji (rozpoznanie kolumn jak w Opisach)."""
    from .opisy import _wymagaj_bibliotek, wczytaj_zestawienie
    _wymagaj_bibliotek()                       # openpyxl ładowany w opisach na żądanie
    wynik = []
    for w in wczytaj_zestawienie(xlsx):
        nr = nr_stacji(w.get("plik") or w.get("trafo"))
        if nr is not None:
            wynik.append((w.get("trafo") or str(nr), nr, w.get("miejscowosc") or ""))
    return wynik


def popraw_zestawienie(xlsx: str, nazwy: dict) -> int:
    """Wpisuje poprawione nazwy do kolumny „Miejscowość” zestawienia stacji. Zwraca liczbę zmian."""
    import openpyxl

    from .opisy import _NAGLOWKI, _ascii
    wb = openpyxl.load_workbook(xlsx)
    ws = wb[wb.sheetnames[0]]
    naglowek = [_ascii("" if c.value is None else str(c.value)) for c in ws[1]]
    k_m = next((i for i, h in enumerate(naglowek) if h in _NAGLOWKI["miejscowosc"]), None)
    k_nr = next((i for i, h in enumerate(naglowek) if h in _NAGLOWKI["plik"]), None)
    k_tr = next((i for i, h in enumerate(naglowek) if h in _NAGLOWKI["trafo"]), None)
    if k_m is None or (k_nr is None and k_tr is None):
        return 0
    zmian = 0
    for row in ws.iter_rows(min_row=2):
        zrodlo = row[k_nr].value if k_nr is not None else row[k_tr].value
        nr = nr_stacji(zrodlo)
        if nr in nazwy and row[k_m].value != nazwy[nr]:
            row[k_m].value = nazwy[nr]
            zmian += 1
    if zmian:
        wb.save(xlsx)
    return zmian
