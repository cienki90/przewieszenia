#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Moduł „Zakładki” - komplet stron-przekładek (zakładek) dla całego projektu.

Dla każdej stacji trafo z zestawienia powstaje jedna strona z trzema wyśrodkowanymi
liniami (dużą czcionką):

    ST 05-0112
    Kiczki Pierwsze
    71/RE-5/KŚ/2026

tzn.:
    1. numer stacji trafo (z prefiksem „ST ”),
    2. miejscowość,
    3. numer warunków technicznych (WT).

Dane pobierane są z pliku zestawienia - najczęściej z „zestawienie_stacji.xlsx”
generowanego przez moduł „Rysunki stałe” (orientacje). Rozpoznawanie kolumn jest
wspólne z modułem „Opisy” (ten sam parser wczytaj_zestawienie z opisy.py), więc
działa też z zestawieniem w formacie „domtel” (Plik | trafo | miejscowosc | | wt).
Numer WT - jeden wspólny - można podać ręcznie, gdy zestawienie go nie zawiera.

Wynik powstaje w JEDNYM z dwóch formatów (przełącznik):
  * DOCX - jeden dokument .docx, każda stacja na osobnej stronie (wymaga python-docx),
  * PDF  - jeden plik .pdf, każda stacja na osobnej stronie (czysty Python, bez bibliotek).

Użycie (CLI):
    python -m generator.zakladki "zestawienie_stacji.xlsx" --wt 71/RE-5/KŚ/2026
    python -m generator.zakladki "zestawienie_stacji.xlsx" --format pdf -o zakladki.pdf
"""
from __future__ import annotations

import argparse
import os
import struct
import sys

try:  # część pakietu (python -m generator.zakladki / okno)
    from . import opisy as _opisy
except ImportError:  # bezpośrednie uruchomienie: python generator/zakladki.py
    import os as _os
    import sys as _sys
    _sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
    import opisy as _opisy

FORMATY = ("docx", "pdf")


# ---------------------------------------------------------------------------
# Przygotowanie danych
# ---------------------------------------------------------------------------
def _norm_stacja(trafo: str) -> str:
    """Numer stacji z prefiksem „ST ” (bez podwajania, gdy już jest)."""
    s = (trafo or "").strip()
    if not s:
        return ""
    if s.upper().startswith("ST "):
        return "ST " + s[3:].strip()
    return "ST " + s


def przygotuj_strony(zestawienie: str, arkusz: str = "Arkusz1", wt: str = ""):
    """Zwraca listę krotek (nr_stacji, miejscowosc, wt) - po jednej na stację.

    Dane czytane są przez wspólny parser z modułu „Opisy” (rozpoznaje kolumny po
    nagłówkach - format „domtel” i „Orientacje”). WT bez wartości w zestawieniu
    uzupełniany jest wartością `wt` (ręcznie wpisany numer wspólny)."""
    _opisy._wymagaj_bibliotek()  # potrzebny openpyxl do odczytu .xlsx
    wiersze = _opisy.wczytaj_zestawienie(zestawienie, arkusz, wt_domyslne=wt)
    strony = []
    for w in wiersze:
        strony.append((_norm_stacja(w.get("trafo") or w.get("plik") or ""),
                       (w.get("miejscowosc") or "").strip(),
                       (w.get("wt") or "").strip()))
    return strony


# ---------------------------------------------------------------------------
# Rozmiar strony zakładki (koperta/przekładka Com-10): 104.7 x 241.3 mm.
# Tekst jest wyrównany do GÓRY strony, wyśrodkowany poziomo, czcionką pogrubioną 12 pt.
# ---------------------------------------------------------------------------
_COM10_W_MM = 104.7
_COM10_H_MM = 241.3
_MARGINES_GORA_MM = 15.0          # tekst zaczyna się 1,5 cm od górnej krawędzi
_ROZMIAR_PT = 12                  # pogrubiona czcionka 12 pt


# ---------------------------------------------------------------------------
# Wynik: DOCX
# ---------------------------------------------------------------------------
def _zapisz_docx(strony, wyjscie: str, rozmiar_pt: int = _ROZMIAR_PT):
    """Jeden .docx: każda stacja na osobnej stronie (Com-10), 3 wyśrodkowane linie
    pogrubioną czcionką, wyrównane do góry strony."""
    try:
        import docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Mm, Pt
    except ImportError as ex:
        raise ImportError("Format DOCX wymaga biblioteki python-docx "
                          "(pip install python-docx)") from ex

    doc = docx.Document()
    sek = doc.sections[0]
    # rozmiar strony Com-10 (104.7 x 241.3 mm) + wąskie marginesy, tekst przy górze
    sek.page_width = Mm(_COM10_W_MM)
    sek.page_height = Mm(_COM10_H_MM)
    sek.top_margin = Mm(_MARGINES_GORA_MM)
    sek.bottom_margin = Mm(5)
    sek.left_margin = Mm(5)
    sek.right_margin = Mm(5)

    for i, (stacja, miejscowosc, wt) in enumerate(strony):
        if i > 0:
            doc.add_page_break()
        for tekst in (stacja, miejscowosc, wt):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf = p.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            run = p.add_run(tekst)
            run.bold = True
            run.font.size = Pt(rozmiar_pt)
    doc.save(wyjscie)
    return wyjscie


# ---------------------------------------------------------------------------
# Wynik: PDF (czysty Python - bez bibliotek zewnętrznych)
# ---------------------------------------------------------------------------
# PDF z OSADZONĄ czcionką TrueType (DejaVu Sans Bold). Osadzenie prawdziwej czcionki
# daje poprawne polskie znaki (ą ć ę ł ń ś ź ż …) w KAŻDYM czytniku - także w Adobe
# Acrobat Reader (który bywa rygorystyczny wobec glifów bazowych czcionek z /Differences
# i potrafił „zjadać” sąsiednie znaki, np. ukośnik po Ś).
#
# Składamy czcionkę złożoną Type0 / CIDFontType2 z kodowaniem Identity-H:
#   * każdy znak zapisujemy jako 2-bajtowy numer glifu (GID) czcionki,
#   * tablica /W podaje szerokości glifów, a /ToUnicode zapewnia kopiowanie/wyszukiwanie,
#   * cały plik .ttf jest osadzony (FontFile2) - plik jest samowystarczalny.
#
# Strona zakładki ma rozmiar Com-10 (104.7 x 241.3 mm), tekst wyrównany do GÓRY.
_MM_PT = 72.0 / 25.4                        # mm -> punkty PDF
_STR_W = _COM10_W_MM * _MM_PT               # szerokość strony [pt]
_STR_H = _COM10_H_MM * _MM_PT               # wysokość strony [pt]
_GORA_PT = _MARGINES_GORA_MM * _MM_PT       # margines górny [pt]


def _czcionka_ttf_sciezka():
    """Ścieżka do dołączonej czcionki (DejaVu Sans Bold) - obok tego modułu lub w zasobach."""
    import os as _os
    tu = _os.path.dirname(_os.path.abspath(__file__))
    kandydaci = [_os.path.join(tu, "fonty", "DejaVuSans-Bold.ttf")]
    meipass = getattr(sys, "_MEIPASS", None)  # zasoby wbudowane w .exe (PyInstaller)
    if meipass:
        kandydaci.append(_os.path.join(meipass, "fonty", "DejaVuSans-Bold.ttf"))
        kandydaci.append(_os.path.join(meipass, "DejaVuSans-Bold.ttf"))
    for p in kandydaci:
        if _os.path.exists(p):
            return p
    raise FileNotFoundError("Brak dołączonej czcionki DejaVuSans-Bold.ttf "
                            "(oczekiwana w generator/fonty/).")


def _tekst_identity(ttf, s: str) -> bytes:
    """Zamienia tekst na 2-bajtowe numery glifów (GID) dla kodowania Identity-H."""
    out = bytearray()
    for ch in s:
        g = ttf.gid(ch) or ttf.gid("?")
        out += struct.pack(">H", g & 0xFFFF)
    return bytes(out)


def _hex_identity(ttf, s: str) -> bytes:
    """Łańcuch glifów jako ciąg szesnastkowy PDF <...> (bezpieczny dla dowolnych bajtów)."""
    return b"<" + _tekst_identity(ttf, s).hex().encode("ascii") + b">"


def _szerokosc_tekstu(ttf, s: str, rozmiar: float) -> float:
    """Szerokość tekstu [pt] z realnych metryk osadzonej czcionki (do wyśrodkowania)."""
    return sum(ttf.advance_1000(ch) for ch in s) / 1000.0 * rozmiar


def _strumien_strony(ttf, stacja: str, miejscowosc: str, wt: str,
                     rozmiar: float = _ROZMIAR_PT) -> bytes:
    """Strumień rysujący 3 wyśrodkowane linie przy GÓRZE strony Com-10 (Identity-H)."""
    linie = [stacja, miejscowosc, wt]
    interlinia = rozmiar * 1.2                       # ciasna interlinia (jak we wzorcu)
    y0 = _STR_H - _GORA_PT - rozmiar                 # pierwsza linia tuż pod górnym marginesem
    polecenia = bytearray(b"BT\n")
    for i, tekst in enumerate(linie):
        szer = _szerokosc_tekstu(ttf, tekst, rozmiar)
        x = (_STR_W - szer) / 2
        y = y0 - i * interlinia
        polecenia += b"/F1 %d Tf\n" % int(rozmiar)
        polecenia += b"1 0 0 1 %.2f %.2f Tm\n" % (x, y)
        polecenia += _hex_identity(ttf, tekst) + b" Tj\n"
    polecenia += b"ET\n"
    return bytes(polecenia)


def _tablica_W(ttf, znaki) -> bytes:
    """Tablica /W (szerokości glifów) dla użytych znaków: [gid [w] gid [w] ...]."""
    pary = {}
    for ch in znaki:
        g = ttf.gid(ch)
        if g:
            pary[g] = ttf.advance_1000(ch)
    czesci = [b"%d [%d]" % (g, w) for g, w in sorted(pary.items())]
    return b" ".join(czesci)


def _tourn_cmap(ttf, znaki) -> bytes:
    """CMap /ToUnicode: GID -> Unicode (kopiowanie i wyszukiwanie tekstu w czytniku)."""
    pary = {}
    for ch in znaki:
        g = ttf.gid(ch)
        if g:
            pary[g] = ord(ch)
    linie = [b"<%04X> <%04X>" % (g, u) for g, u in sorted(pary.items())]
    body = (b"/CIDInit /ProcSet findresource begin\n12 dict begin\nbegincmap\n"
            b"/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def\n"
            b"/CMapName /Adobe-Identity-UCS def\n/CMapType 2 def\n"
            b"1 begincodespacerange <0000> <FFFF> endcodespacerange\n"
            + b"%d beginbfchar\n" % len(linie) + b"\n".join(linie)
            + b"\nendbfchar\nendcmap CMapName currentdict /CMap defineresource pop\nend\nend")
    return body


def _zapisz_pdf(strony, wyjscie: str, rozmiar_pt: float = _ROZMIAR_PT):
    """Jeden .pdf z OSADZONĄ czcionką TrueType: każda stacja na osobnej stronie Com-10."""
    if not strony:
        raise ValueError("Brak danych do zapisania w PDF.")

    with open(_czcionka_ttf_sciezka(), "rb") as f:
        ttf_bytes = f.read()
    try:
        from . import _ttf as _ttfmod
    except ImportError:
        import _ttf as _ttfmod
    ttf = _ttfmod.TTF(ttf_bytes)

    # zbiór wszystkich użytych znaków (do /W i /ToUnicode)
    znaki = set()
    for stacja, miejscowosc, wt in strony:
        znaki.update(stacja + miejscowosc + wt)

    s = 1000.0 / ttf.unitsPerEm
    bbox = [round(v * s) for v in ttf.bbox]

    # numeracja obiektów:
    # 1 katalog, 2 drzewo stron, 3 Type0, 4 CIDFont, 5 FontDescriptor, 6 FontFile2,
    # 7 ToUnicode; strony i treści: 8,9 (strona 1), 10,11 (strona 2), ...
    n_stron = len(strony)
    kids_num = [8 + 2 * i for i in range(n_stron)]
    obiekty = []

    obiekty.append(b"<< /Type /Catalog /Pages 2 0 R >>")                              # 1
    kids = b" ".join(b"%d 0 R" % k for k in kids_num)
    obiekty.append(b"<< /Type /Pages /Count %d /Kids [%s] >>" % (n_stron, kids))      # 2
    # 3: czcionka złożona Type0 (Identity-H) odwołująca się do CIDFont (4) i ToUnicode (7)
    obiekty.append(b"<< /Type /Font /Subtype /Type0 /BaseFont /DejaVuSans-Bold "
                   b"/Encoding /Identity-H /DescendantFonts [4 0 R] /ToUnicode 7 0 R >>")  # 3
    # 4: CIDFontType2 z tablicą szerokości /W i mapą CID->GID /Identity
    obiekty.append(b"<< /Type /Font /Subtype /CIDFontType2 /BaseFont /DejaVuSans-Bold "
                   b"/CIDSystemInfo << /Registry (Adobe) /Ordering (Identity) /Supplement 0 >> "
                   b"/FontDescriptor 5 0 R /CIDToGIDMap /Identity /DW 500 /W [%s] >>"
                   % _tablica_W(ttf, znaki))                                           # 4
    # 5: opis czcionki (symbolic, bbox z pliku .ttf) odwołujący się do FontFile2 (6)
    obiekty.append(b"<< /Type /FontDescriptor /FontName /DejaVuSans-Bold /Flags 4 "
                   b"/FontBBox [%d %d %d %d] /ItalicAngle 0 /Ascent %d /Descent %d "
                   b"/CapHeight %d /StemV 140 /FontFile2 6 0 R >>"
                   % (bbox[0], bbox[1], bbox[2], bbox[3], bbox[3], bbox[1], bbox[3]))  # 5
    # 6: osadzony plik czcionki (cały .ttf)
    obiekty.append(b"<< /Length %d /Length1 %d >>\nstream\n" % (len(ttf_bytes), len(ttf_bytes))
                   + ttf_bytes + b"\nendstream")                                       # 6
    # 7: ToUnicode CMap
    tu = _tourn_cmap(ttf, znaki)
    obiekty.append(b"<< /Length %d >>\nstream\n" % len(tu) + tu + b"\nendstream")      # 7

    for i, (stacja, miejscowosc, wt) in enumerate(strony):
        tresc = _strumien_strony(ttf, stacja, miejscowosc, wt, rozmiar_pt)
        obj_tresc = kids_num[i] + 1
        strona = (b"<< /Type /Page /Parent 2 0 R "
                  b"/MediaBox [0 0 %.2f %.2f] "
                  b"/Resources << /Font << /F1 3 0 R >> >> "
                  b"/Contents %d 0 R >>" % (_STR_W, _STR_H, obj_tresc))
        obiekty.append(strona)
        obiekty.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(tresc), tresc))

    # złożenie pliku z tablicą xref
    buf = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsety = [0] * (len(obiekty) + 1)
    for idx, body in enumerate(obiekty, start=1):
        offsety[idx] = len(buf)
        buf += b"%d 0 obj\n" % idx + body + b"\nendobj\n"
    xref_pos = len(buf)
    buf += b"xref\n0 %d\n" % (len(obiekty) + 1)
    buf += b"0000000000 65535 f \n"
    for idx in range(1, len(obiekty) + 1):
        buf += b"%010d 00000 n \n" % offsety[idx]
    buf += (b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
            % (len(obiekty) + 1, xref_pos))

    with open(wyjscie, "wb") as f:
        f.write(buf)
    return wyjscie


# ---------------------------------------------------------------------------
# Funkcja generująca (wywoływalna z CLI i z pipeline/GUI)
# ---------------------------------------------------------------------------
def generuj(zestawienie, wyjscie=None, format="docx", arkusz="Arkusz1", wt="", log=None):
    """Tworzy plik zakładek (DOCX albo PDF) z zestawienia.

    Parametry:
        zestawienie - ścieżka do .xlsx (np. „zestawienie_stacji.xlsx” z Rysunków stałych),
        wyjscie     - plik wynikowy; gdy None, obok zestawienia (zakladki.docx/.pdf),
        format      - „docx” albo „pdf” (przełącznik formatu wyniku),
        arkusz      - nazwa arkusza z danymi (domyślnie „Arkusz1”),
        wt          - numer warunków (WT) wspólny, gdy zestawienie go nie zawiera,
        log         - funkcja log(msg) (domyślnie print).

    Zwraca ścieżkę utworzonego pliku."""
    if log is None:
        log = print
    fmt = (format or "docx").strip().lower()
    if fmt not in FORMATY:
        raise ValueError(f"Nieznany format „{format}” - dozwolone: {', '.join(FORMATY)}.")
    if not os.path.isfile(zestawienie):
        raise FileNotFoundError(f"Nie znaleziono zestawienia: {zestawienie}")

    strony = przygotuj_strony(zestawienie, arkusz, wt=wt)
    if not strony:
        raise ValueError("Zestawienie nie zawiera danych (żadnej stacji).")

    if wyjscie is None:
        baza = os.path.join(os.path.dirname(os.path.abspath(zestawienie)), "zakladki")
        wyjscie = baza + "." + fmt
    else:
        # dopilnuj rozszerzenia zgodnego z formatem
        rdzen, ext = os.path.splitext(wyjscie)
        if ext.lower() != "." + fmt:
            wyjscie = rdzen + "." + fmt
    os.makedirs(os.path.dirname(os.path.abspath(wyjscie)), exist_ok=True)

    braki_wt = [s[0] for s in strony if not s[2]]
    if braki_wt:
        log(f"Uwaga: brak numeru WT dla {len(braki_wt)} stacji (np. {braki_wt[:3]}) - "
            f"podaj numer warunków (--wt) lub uzupełnij kolumnę „wt” w zestawieniu.")

    log(f"Znaleziono {len(strony)} stacji w zestawieniu. Format: {fmt.upper()}.")
    if fmt == "docx":
        _zapisz_docx(strony, wyjscie)
    else:
        _zapisz_pdf(strony, wyjscie)
    for stacja, miejscowosc, w in strony:
        log(f"  [OK] {stacja:12s} {miejscowosc:20s} WT={w}")
    log(f"Gotowe: {wyjscie}")
    return wyjscie


# ---------------------------------------------------------------------------
# Main (CLI)
# ---------------------------------------------------------------------------
def main(argv=None):
    for strumien in (sys.stdout, sys.stderr):  # polskie znaki w konsoli Windows
        try:
            strumien.reconfigure(encoding="utf-8", errors="replace")
        except Exception:  # noqa: BLE001
            pass
    ap = argparse.ArgumentParser(
        description="Zakładki (przekładki): dla każdej stacji trafo z zestawienia osobna strona "
                    "(nr stacji / miejscowość / WT). Wynik .docx albo .pdf.")
    ap.add_argument("zestawienie", help="plik .xlsx z danymi (np. „zestawienie_stacji.xlsx” z Rysunków stałych)")
    ap.add_argument("-o", "--wyjscie", default=None,
                    help="plik wynikowy (domyślnie zakladki.<format> obok zestawienia)")
    ap.add_argument("--format", choices=FORMATY, default="docx",
                    help="format wyniku: docx (domyślnie) albo pdf")
    ap.add_argument("--arkusz", default="Arkusz1", help="nazwa arkusza z danymi (domyślnie Arkusz1)")
    ap.add_argument("--wt", default="",
                    help="numer warunków technicznych (WT) - jeden wspólny, gdy zestawienie go nie zawiera")
    a = ap.parse_args(argv)
    try:
        generuj(a.zestawienie, a.wyjscie, format=a.format, arkusz=a.arkusz, wt=a.wt)
    except Exception as ex:  # noqa: BLE001
        sys.exit(str(ex))
    return 0


if __name__ == "__main__":
    sys.exit(main())
