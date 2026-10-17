# -*- coding: utf-8 -*-
"""
Moduł obliczenia2word.

Generuje do dokumentu .docx tabelę "Zestawienie zbiorcze słupów"
(obliczenia nośności słupów) na podstawie arkusza projektu z pliku
"Obliczenia - zestawienie pw.xlsx".

Każdy projekt ma w skoroszycie Excela własny arkusz nazwany numerem
(np. '112', '115', ...). Arkusz zawiera jeden lub kilka obwodów; każdy
obwód to blok:

    Sieć napowietrzna nN | ... | obw nr: | N | | STACJA TRAFO: 05- | 0112
    Lp | | Nr słupa | Nośność [daN] | Obciążenie [daN] | Typ i rodzaj słupa |
       Rodzaj kabla ... | Osprzęt | Odległość [m]
    <wiersze danych słupów>
    Całkowita długość wykorzystanej linii napowietrznej nN: | ... | <X m>
    Całkowita ilość wykorzystanych stanowisk słupowych: | ... | <N>

Na końcu arkusza znajdują się wiersze zbiorcze (scalone na całą szerokość):
    WSZYSTKIE STANOWISKA ... / UWAGA! NIE WSZYSTKIE ...
    CAŁKOWITA DŁUGOŚĆ ...
    CAŁKOWITA ILOŚĆ ...
oraz wiersze z projektantami.

Tabela w .docx odwzorowuje arkusz 1:1, przy czym kolumna "Obciążenie [daN]"
jest zaokrąglana do 2 miejsc po przecinku.

Formatowanie (czcionki Lato / Lato Light, obramowania, scalenia) jest
pobierane z pliku-wzorca "wzorce/wzor_tabeli_obliczen.docx" przez klonowanie
odpowiednich wierszy – dzięki temu wygląd tabeli jest identyczny z wzorcem.
"""

import copy
import os

from docx import Document
from docx.oxml.ns import qn

def _wzorzec_domyslny():
    """Ścieżka do wzorca tabeli obliczeń - z zasobów pakietu / .exe, a w razie czego obok modułu."""
    try:
        from . import zasoby
        p = zasoby.wzor_tabeli_obliczen_docx()
        if p:
            return p
    except Exception:  # noqa: BLE001
        pass
    return os.path.join(os.path.dirname(os.path.abspath(__file__)), "wzorce", "wzor_tabeli_obliczen.docx")


WZORZEC_DOMYSLNY = _wzorzec_domyslny()


# ---------------------------------------------------------------------------
# Pomocnicze: tekst / typy wartości
# ---------------------------------------------------------------------------
def _s(v):
    """Zamiana wartości komórki Excela na tekst (bez zbędnego '.0'), z przycięciem."""
    return _s_raw(v).strip()


def _s_raw(v):
    """Jak _s, ale BEZ przycinania spacji (wierne odwzorowanie treści komórki)."""
    if v is None:
        return ""
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return str(v)
    return str(v)


def _fmt_obciazenie(v):
    """Obciążenie [daN] -> zawsze 2 miejsca po przecinku (np. 928.50)."""
    if v is None or _s(v) == "":
        return ""
    try:
        return f"{float(v):.2f}"
    except (TypeError, ValueError):
        return _s(v)


# ---------------------------------------------------------------------------
# Parsowanie arkusza projektu
# ---------------------------------------------------------------------------
def parsuj_arkusz(ws):
    """
    Zwraca listę "wierszy logicznych" opisujących zawartość tabeli:
      ('naglowek_obw', {'obw': '1', 'trafo_prefix': 'STACJA TRAFO: 05-', 'trafo': '0112'})
      ('naglowek_kol', [9 nagłówków])
      ('dana', [9 wartości])
      ('suma_dlugosc', '442 m')
      ('suma_slupy', '11')
      ('uwaga', 'WSZYSTKIE STANOWISKA ...')          # wiersz scalony na całość
      ('pusty', None)
      ('projektant', ['Projektant:', 'inż. ...', 'PDL/...'])
    """
    wiersze = []
    for row in ws.iter_rows(values_only=True):
        vals = list(row)
        a = _s(vals[0]) if vals else ""

        if not any(_s(v) for v in vals):
            continue

        # nagłówek obwodu
        if a.startswith("Sieć napowietrzna nN"):
            obw = ""
            trafo_prefix = ""
            trafo = ""
            # obw nr: wartość – szukamy etykiety "obw nr:" i bierzemy następną komórkę
            for i, v in enumerate(vals):
                if _s(v).lower().startswith("obw nr"):
                    if i + 1 < len(vals):
                        obw = _s(vals[i + 1])
                if _s(v).upper().startswith("STACJA TRAFO"):
                    trafo_prefix = _s(v)
                    if i + 1 < len(vals):
                        trafo = _s(vals[i + 1])
            wiersze.append(("naglowek_obw",
                            {"obw": obw, "trafo_prefix": trafo_prefix, "trafo": trafo}))
            continue

        # nagłówek kolumn
        if a == "Lp":
            wiersze.append(("naglowek_kol", [_s(v) for v in vals[:9]]))
            continue

        # podsumowania obwodu
        if a.startswith("Całkowita długość"):
            wiersze.append(("suma_dlugosc", _s(vals[-1])))
            continue
        if a.startswith("Całkowita ilość"):
            wiersze.append(("suma_slupy", _s(vals[-1])))
            continue

        # wiersze zbiorcze (scalone na całą szerokość)
        if a.startswith("WSZYSTKIE STANOWISKA") or a.startswith("UWAGA!") \
                or a.startswith("CAŁKOWITA DŁUGOŚĆ") or a.startswith("CAŁKOWITA ILOŚĆ"):
            wiersze.append(("uwaga", a))
            continue

        # projektanci
        if a.startswith("Projektant"):
            # imię i nr uprawnień bywają w dalszych kolumnach
            niepuste = [_s(v) for v in vals if _s(v)]
            wiersze.append(("projektant", niepuste))
            continue

        # pozostałe – wiersz danych słupa (9 kolumn)
        # wartości wierne (bez przycinania spacji), poza kolumną obciążenia
        dana = [_s_raw(v) for v in vals[:9]]
        if len(dana) > 4:
            dana[4] = _fmt_obciazenie(vals[4] if len(vals) > 4 else None)
        wiersze.append(("dana", dana))

    # Pusty wiersz-separator przed sekcją projektantów (jak w layoutcie wzorca)
    idx_proj = next((i for i, (t, _) in enumerate(wiersze) if t == "projektant"), None)
    if idx_proj is not None and (idx_proj == 0 or wiersze[idx_proj - 1][0] != "pusty"):
        wiersze.insert(idx_proj, ("pusty", None))

    return wiersze


# ---------------------------------------------------------------------------
# Wzorce wierszy z pliku-wzorca
# ---------------------------------------------------------------------------
class WzorceWierszy:
    """
    Wyodrębnia z tabeli obliczeń pliku-wzorca po jednym <w:tr> każdego typu,
    aby móc je klonować przy budowie nowej tabeli.
    """
    def __init__(self, sciezka_wzorca=WZORZEC_DOMYSLNY):
        doc = Document(sciezka_wzorca)
        tbl = _znajdz_tabele_obliczen(doc)
        if tbl is None:
            raise RuntimeError("Nie znaleziono tabeli obliczeń w pliku-wzorcu.")
        self.tbl_wzor = tbl._tbl
        trs = self.tbl_wzor.findall(qn("w:tr"))

        self.tblPr = copy.deepcopy(self.tbl_wzor.find(qn("w:tblPr")))
        self.tblGrid = copy.deepcopy(self.tbl_wzor.find(qn("w:tblGrid")))

        # klasyfikacja wierszy wzorca
        self.wz = {}
        for tr in trs:
            tcs = tr.findall(qn("w:tc"))
            txt0 = _tekst_tc(tcs[0]) if tcs else ""
            ntc = len(tcs)
            if txt0.startswith("Sieć napowietrzna nN") and "naglowek_obw" not in self.wz:
                self.wz["naglowek_obw"] = tr
            elif txt0 == "Lp" and "naglowek_kol" not in self.wz:
                self.wz["naglowek_kol"] = tr
            elif txt0.startswith("Całkowita długość") and "suma_dlugosc" not in self.wz:
                self.wz["suma_dlugosc"] = tr
            elif txt0.startswith("Całkowita ilość") and "suma_slupy" not in self.wz:
                self.wz["suma_slupy"] = tr
            elif ntc == 1 and "uwaga" not in self.wz:
                self.wz["uwaga"] = tr
            elif txt0.startswith("Projektant") and "projektant" not in self.wz:
                self.wz["projektant"] = tr
            elif ntc == 9 and txt0 == "" and "pusty" not in self.wz:
                self.wz["pusty"] = tr
            elif ntc == 9 and "dana" not in self.wz and txt0 not in ("Lp",):
                # pierwszy "zwykły" 9-kolumnowy wiersz danych
                self.wz["dana"] = tr

    def klon(self, typ):
        return copy.deepcopy(self.wz[typ])


def _tekst_tc(tc):
    return "".join(n.text or "" for n in tc.iter(qn("w:t")))


def _znajdz_tabele_obliczen(doc):
    """Tabela obliczeń to ta, której pierwszy wiersz zaczyna się od 'Sieć napowietrzna nN'."""
    for t in doc.tables:
        if t.rows and _tekst_tc(t.rows[0].cells[0]._tc).startswith("Sieć napowietrzna nN"):
            return t
    return None


# ---------------------------------------------------------------------------
# Ustawianie tekstów w sklonowanym wierszu
# ---------------------------------------------------------------------------
def _ustaw_teksty_tc(tc, tekst):
    """
    Podmienia tekst w komórce: zostawia pierwszy run (z jego formatowaniem),
    wpisuje do niego 'tekst', usuwa pozostałe runy w pierwszym akapicie
    i czyści dalsze akapity.
    """
    paras = tc.findall(qn("w:p"))
    if not paras:
        return
    p0 = paras[0]
    runs = p0.findall(qn("w:r"))
    if runs:
        # ustaw tekst w pierwszym runie
        _ustaw_tekst_run(runs[0], tekst)
        # usuń pozostałe runy
        for r in runs[1:]:
            p0.remove(r)
    else:
        # brak runów – stwórz prosty run bez formatowania
        from docx.oxml import OxmlElement
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = tekst
        r.append(t)
        p0.append(r)
    # wyczyść dalsze akapity
    for p in paras[1:]:
        for r in p.findall(qn("w:r")):
            p.remove(r)


def _ustaw_tekst_run(r, tekst):
    for t in r.findall(qn("w:t")):
        r.remove(t)
    from docx.oxml import OxmlElement
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = tekst
    r.append(t)


def _komorki(tr):
    return tr.findall(qn("w:tc"))


def _wypelnij_wiersz(tr, typ, dane):
    """Wpisuje dane do sklonowanego wiersza zgodnie z jego typem."""
    tcs = _komorki(tr)

    if typ == "naglowek_obw":
        # układ komórek: [Sieć... (span4)], [obw nr:], [N], [nbsp], [STACJA TRAFO: 05-], [trafo]
        if len(tcs) >= 6:
            _ustaw_teksty_tc(tcs[2], dane.get("obw", ""))
            if dane.get("trafo_prefix"):
                _ustaw_teksty_tc(tcs[4], dane["trafo_prefix"])
            _ustaw_teksty_tc(tcs[5], dane.get("trafo", ""))

    elif typ == "naglowek_kol":
        # wzorzec ma już poprawne nagłówki – nie ruszamy (identyczne)
        pass

    elif typ == "dana":
        # 9 komórek = 9 wartości
        for i, tc in enumerate(tcs):
            if i < len(dane):
                _ustaw_teksty_tc(tc, dane[i])
            else:
                _ustaw_teksty_tc(tc, "")

    elif typ in ("suma_dlugosc", "suma_slupy"):
        # komórki: [etykieta (span8)], [wartość]
        if len(tcs) >= 2:
            _ustaw_teksty_tc(tcs[-1], dane)

    elif typ == "uwaga":
        # jedna komórka scalona na całość
        if tcs:
            _ustaw_teksty_tc(tcs[0], dane)

    elif typ == "projektant":
        # komórki: [Projektant: (span4)], [imię (span2)], [nr upr], [nbsp (span2)]
        if dane:
            if len(tcs) >= 1:
                _ustaw_teksty_tc(tcs[0], dane[0] if len(dane) > 0 else "")
            if len(tcs) >= 2 and len(dane) > 1:
                _ustaw_teksty_tc(tcs[1], dane[1])
            if len(tcs) >= 3 and len(dane) > 2:
                _ustaw_teksty_tc(tcs[2], dane[2])


# ---------------------------------------------------------------------------
# Budowa tabeli
# ---------------------------------------------------------------------------
def zbuduj_tabele(wiersze_logiczne, wzorce):
    """
    Tworzy nowy element <w:tbl> z wierszami zbudowanymi z 'wiersze_logiczne'
    (wynik parsuj_arkusz) przy użyciu wzorców.
    """
    from docx.oxml import OxmlElement
    tbl = OxmlElement("w:tbl")
    tbl.append(copy.deepcopy(wzorce.tblPr))
    tbl.append(copy.deepcopy(wzorce.tblGrid))

    for typ, dane in wiersze_logiczne:
        if typ not in wzorce.wz:
            # brak wzorca danego typu – pomijamy bezpiecznie
            continue
        tr = wzorce.klon(typ)
        _wypelnij_wiersz(tr, typ, dane)
        tbl.append(tr)

    return tbl


# ---------------------------------------------------------------------------
# Wstawienie tabeli do dokumentu
# ---------------------------------------------------------------------------
def wstaw_tabele_obliczen(doc, ws, wzorce=None, naglowek="Zestawienie zbiorcze słupów"):
    """
    Wstawia do dokumentu 'doc' tabelę obliczeń zbudowaną z arkusza 'ws'.
    Tabela trafia bezpośrednio po akapicie nagłówka sekcji 'naglowek'
    (pomijając akapit spisu treści). Jeśli dokument ma już taką tabelę
    (pierwszy wiersz 'Sieć napowietrzna nN'), zostaje ona zastąpiona.

    Zwraca liczbę zbudowanych wierszy tabeli.
    """
    if wzorce is None:
        wzorce = WzorceWierszy()

    wiersze = parsuj_arkusz(ws)
    nowa_tbl = zbuduj_tabele(wiersze, wzorce)

    # 1) jeśli istnieje już tabela obliczeń – podmień ją w miejscu
    istn = _znajdz_tabele_obliczen(doc)
    if istn is not None:
        istn._tbl.addnext(nowa_tbl)
        istn._tbl.getparent().remove(istn._tbl)
        return len(nowa_tbl.findall(qn("w:tr")))

    # 2) w przeciwnym razie wstaw po nagłówku sekcji (poza spisem treści)
    cel = _znajdz_akapit_naglowka(doc, naglowek)
    if cel is None:
        # fallback: na koniec dokumentu
        doc.element.body.append(nowa_tbl)
    else:
        cel.addnext(nowa_tbl)

    return len(nowa_tbl.findall(qn("w:tr")))


def _znajdz_akapit_naglowka(doc, naglowek):
    """
    Znajduje element <w:p> będący nagłówkiem sekcji (nie wpisem spisu treści).
    Wpisy spisu treści zawierają tabulatory i numer strony – pomijamy je.
    """
    trafienia = []
    for p in doc.paragraphs:
        if p.text.strip() == naglowek:
            trafienia.append(p._p)
    if not trafienia:
        return None
    # ostatnie trafienie to zwykle właściwy nagłówek sekcji (spis treści jest wcześniej)
    return trafienia[-1]
